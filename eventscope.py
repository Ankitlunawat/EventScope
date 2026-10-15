"""Extract event-calendar records for conference research and planning."""
import argparse,csv,gzip,html,json,urllib.request
from datetime import datetime,timezone,date
from pathlib import Path
from urllib.parse import urljoin
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
ROOT=Path(__file__).resolve().parent
SOURCE='https://www.python.org/events/python-events/'
def parse(source,base=SOURCE):
 rows=[];seen=set();issues=[]
 for item in BeautifulSoup(source,'html.parser').select('ul.list-recent-events li'):
  a=item.select_one('.event-title a');t=item.select_one('time[datetime]');loc=item.select_one('.event-location')
  if a is None or t is None:issues.append('Skipped listing missing title or start time');continue
  url=urljoin(base,a.get('href',''));title=a.get_text(' ',strip=True)
  if not url.startswith('https://www.python.org/events/') or not title:issues.append('Skipped invalid title or event URL');continue
  try:
   start=datetime.fromisoformat(t['datetime'].replace('Z','+00:00'))
   if start.tzinfo is None:raise ValueError('timezone missing')
  except ValueError:issues.append('Skipped invalid or timezone-free date');continue
  if url in seen:continue
  seen.add(url)
  rows.append({'event':title,'start_iso':start.isoformat(),'start_date':start.date().isoformat(),'source_timezone':start.strftime('%z'),'location':loc.get_text(' ',strip=True) if loc else 'Not listed','event_url':url,'source_url':base})
 return sorted(rows,key=lambda r:(r['start_iso'],r['event'])),issues

def export(rows,out):
 if not rows:raise ValueError('No valid events; inspect saved HTML and selectors')
 out.mkdir(parents=True,exist_ok=True)
 with (out/'events.csv').open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
 (out/'events.json').write_text(json.dumps(rows,indent=2,ensure_ascii=False),encoding='utf-8')
 wb=Workbook();ws=wb.active;ws.title='Events';ws.append(list(rows[0]))
 for row in rows:ws.append(list(row.values()))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='102B3E')
 for col,width in zip('ABCDEFG',[42,30,18,20,45,72,65]):ws.column_dimensions[col].width=width
 wb.save(out/'events.xlsx')
 cards=''.join('<tr>'+''.join('<td>'+html.escape(str(r[k]))+'</td>' for k in ['event','start_date','source_timezone','location'])+'<td><a href="'+html.escape(r['event_url'],quote=True)+'">Source</a></td></tr>' for r in rows)
 (out/'events-report.html').write_text('<!doctype html><html><head><meta charset="utf-8"><title>EventScope | Event research</title><style>body{font:17px system-ui;background:#eef3f8;color:#152f45;margin:50px}h1{font-size:42px}table{width:100%;border-collapse:collapse;background:white}th,td{text-align:left;padding:18px;border-bottom:1px solid #ddd}a{color:#087875}.note{color:#526578}</style></head><body><h1>EventScope</h1><p>'+str(len(rows))+' unique event listings | dates, locations and source links</p><table><tr><th>Event</th><th>Start date</th><th>Source offset</th><th>Location</th><th>Link</th></tr>'+cards+'</table><p class="note">Self-created extraction project. Source: Python.org public event calendar. Snapshot only: verify dates with organizers. Midnight timestamps may reflect source defaults, not session start times.</p></body></html>',encoding='utf-8')

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--refresh',action='store_true');ap.add_argument('--from-date',type=date.fromisoformat);args=ap.parse_args()
 path=ROOT/'source/python-events.html'
 if args.refresh:
  req=urllib.request.Request(SOURCE,headers={'User-Agent':'EventScope-Portfolio/1.0','Accept-Encoding':'identity'})
  with urllib.request.urlopen(req,timeout=30) as r:b=r.read()
  if b[:2]==b'\x1f\x8b':b=gzip.decompress(b)
  path.parent.mkdir(exist_ok=True);path.write_text(b.decode('utf-8'),encoding='utf-8')
 rows,issues=parse(path.read_text(encoding='utf-8'))
 if args.from_date:rows=[r for r in rows if date.fromisoformat(r['start_date'])>=args.from_date]
 export(rows,ROOT/'output')
 report={'processed_at_utc':datetime.now(timezone.utc).isoformat(),'source_url':SOURCE,'source_mode':'fresh fetch' if args.refresh else 'included saved HTML','unique_events':len(rows),'duplicates_in_output':len(rows)-len({r['event_url'] for r in rows}),'issues':issues,'from_date':str(args.from_date) if args.from_date else None,'disclosure':'Self-created project using real public calendar listings; not paid client work. Dates and locations can change. Source timestamps preserved, not inferred session times.'}
 (ROOT/'output/quality_report.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
if __name__=='__main__':main()
